import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
import base64
import urllib.request
import os
import time

def generate_mermaid_image(mermaid_code, filename):
    # コードをUTF-8バイト列にしてURLセーフなBase64に変換し、パディング文字 '=' を除去
    code_bytes = mermaid_code.encode("utf-8")
    base64_bytes = base64.urlsafe_b64encode(code_bytes)
    base64_str = base64_bytes.decode("utf-8").replace("=", "")
    
    url = f"https://mermaid.ink/img/{base64_str}"
    
    images_dir = "g:\\ya-work\\workspace\\tsubuyaki-board\\scratch\\images"
    os.makedirs(images_dir, exist_ok=True)
    img_path = os.path.join(images_dir, f"{filename}.png")
    
    # 接続リトライ処理 (最大3回)
    for attempt in range(3):
        try:
            req = urllib.request.Request(
                url, 
                headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
            )
            with urllib.request.urlopen(req, timeout=20) as response:
                with open(img_path, 'wb') as out_file:
                    out_file.write(response.read())
            print(f"Successfully downloaded image on attempt {attempt + 1}: {img_path}")
            time.sleep(2) # 負荷軽減のため待機時間を2秒に設定
            return img_path
        except Exception as e:
            print(f"Attempt {attempt + 1} failed for {filename}: {e}")
            time.sleep(3) # 失敗時は3秒空けてリトライ
            
    print(f"All attempts failed for {filename}")
    return None

def create_sequence_excel():
    wb = openpyxl.Workbook()
    
    # シート1: 表形式のシーケンス定義
    ws1 = wb.active
    ws1.title = "Sequence Tables"
    ws1.views.sheetView[0].showGridLines = True
    
    # ヘッダー定義
    headers = [
        "機能区分", 
        "機能ID", 
        "機能名", 
        "ステップ", 
        "送信元 (From)", 
        "送信先 (To)", 
        "メッセージ / 処理内容", 
        "物理処理 / 備考 (SQL, バリデーション等)"
    ]
    
    data = [
        # M1: 投稿一覧表示
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 1, "ユーザー (ブラウザ)", "PostController", "GET /posts リクエスト送信", "-"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 2, "PostController", "PostService", "getPosts() 呼び出し", "-"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 3, "PostService", "PostRepository", "findTop50ByDeletedAtIsNullOrderByCreatedAtDesc()", "SQL: SELECT * FROM posts WHERE deleted_at IS NULL ORDER BY created_at DESC (最大50件)"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 4, "PostRepository", "PostService", "投稿リスト返却 (List<Post>)", "-"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 5, "PostService", "PostController", "投稿リスト返却", "-"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 6, "PostController", "Model", "posts 属性にリストを追加", "-"],
        ["MUST", "M1", "投稿一覧表示 (GET /posts)", 7, "PostController", "Thymeleaf (posts/list.html)", "テンプレート描画・HTML返却", "件数0件時は「まだ投稿はありません」表示、新着順に「投稿者」「内容」「投稿日」を出力"],

        # M2: 投稿作成フォーム表示
        ["MUST", "M2", "投稿作成フォーム表示 (GET /posts/new)", 1, "ユーザー (ブラウザ)", "PostController", "GET /posts/new リクエスト送信", "-"],
        ["MUST", "M2", "投稿作成フォーム表示 (GET /posts/new)", 2, "PostController", "Model", "新規の PostForm オブジェクトを設定", "初期状態の空フォーム用"],
        ["MUST", "M2", "投稿作成フォーム表示 (GET /posts/new)", 3, "PostController", "Thymeleaf (posts/form.html)", "テンプレート描画・HTML返却", "投稿者、内容、アバター色、タグの入力欄をバインド表示"],

        # M3: 投稿登録
        ["MUST", "M3", "投稿登録 (POST /posts)", 1, "ユーザー (ブラウザ)", "PostController", "POST /posts リクエスト送信 (Formデータ)", "CSRFトークン検証を含む"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 2, "PostController", "PostController", "バリデーション実行 (@Valid)", "author(1-30文字,空白NG), body(1-280文字,空白NG), color(7文字)"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 3, "PostController", "ClientHashGenerator", "generate(ip, userAgent) 呼び出し", "SHA-256ハッシュの先頭8文字を生成"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 4, "ClientHashGenerator", "PostController", "clientHash 返却", "-"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 5, "PostController", "PostService", "create(form, clientHash) 呼び出し", "異常系(エラーあり)の場合は posts/form 画面を再表示(HTTP 200)"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 6, "PostService", "TagRepository", "findByName(tagName) 呼び出し", "カンマ区切りタグ入力をパースし、既存タグを検索 (再利用)"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 7, "PostService", "TagRepository", "save(tag) 呼び出し", "新規タグの場合のみ SQL: INSERT INTO tags ..."],
        ["MUST", "M3", "投稿登録 (POST /posts)", 8, "PostService", "PostRepository", "save(post) 呼び出し", "SQL: INSERT INTO posts (id, author, body, color, client_hash, created_at) VALUES (...)"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 9, "PostService", "PostController", "保存済み Post 返却", "-"],
        ["MUST", "M3", "投稿登録 (POST /posts)", 10, "PostController", "ユーザー (ブラウザ)", "HTTP 302 Redirect (/posts)", "PRGパターン適用"],

        # M4: 投稿詳細
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 1, "ユーザー (ブラウザ)", "PostController", "GET /posts/{id} リクエスト送信", "-"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 2, "PostController", "PostService", "findById(id) 呼び出し", "-"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 3, "PostService", "PostRepository", "findById(id) 呼び出し", "SQL: SELECT * FROM posts WHERE id = ?"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 4, "PostService", "PostService", "削除状態検証", "deleted_at または purged_at が非NULLなら PostNotFoundException (HTTP 404)"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 5, "PostService", "PostController", "Post エンティティ返却", "-"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 6, "PostController", "PostLikeService", "countByPostId(id) 呼び出し", "いいね総数の取得"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 7, "PostLikeService", "PostLikeRepository", "countByPostId(id)", "SQL: SELECT COUNT(*) FROM post_likes WHERE post_id = ?"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 8, "PostController", "ClientHashGenerator", "generate(ip, userAgent) 呼び出し", "本人検証用の currentHash 生成"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 9, "PostController", "Model", "post, likeCount, currentHash 属性を設定", "-"],
        ["MUST", "M4", "投稿詳細表示 (GET /posts/{id})", 10, "PostController", "Thymeleaf (posts/detail.html)", "テンプレート描画・HTML返却", "投稿内容、アバター色、いいね数、タグ表示、本人にのみ「編集」「削除」ボタン表示"],

        # SHOULD: S1: いいねトグル
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 1, "ユーザー (ブラウザ)", "PostController", "POST /posts/{id}/likes リクエスト送信", "CSRFトークン検証を含む"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 2, "PostController", "ClientHashGenerator", "generate(ip, userAgent) 呼び出し", "同一クライアント判定用ハッシュ生成"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 3, "PostController", "PostLikeService", "toggle(id, clientHash) 呼び出し", "-"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 4, "PostLikeService", "PostRepository", "findById(id) 呼び出し", "投稿の存在確認。なければ 404"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 5, "PostLikeService", "PostLikeRepository", "findByPostIdAndClientHash()", "SQL: SELECT * FROM post_likes WHERE post_id = ? AND client_hash = ?"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 6, "PostLikeService", "PostLikeRepository", "delete(like) / save(like)", "存在すれば物理削除、存在しなければ新規追加(SQL: INSERT / DELETE)"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 7, "PostLikeService", "PostController", "処理完了", "-"],
        ["SHOULD", "S1", "いいねトグル (POST /posts/{id}/likes)", 8, "PostController", "ユーザー (ブラウザ)", "HTTP 302 Redirect (/posts/{id})", "詳細画面へリロード"],

        # SHOULD: S2: 投稿削除
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 1, "ユーザー (ブラウザ)", "PostController", "POST /posts/{id}/delete リクエスト送信", "CSRFトークン検証を含む"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 2, "PostController", "PostService", "deletePost(id) 呼び出し", "-"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 3, "PostService", "PostRepository", "findById(id) 呼び出し", "存在確認。なければ 404"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 4, "PostService", "Post", "delete() 呼び出し", "エンティティの deletedAt フィールドに現在日時を設定"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 5, "PostService", "PostRepository", "save(post) 呼び出し", "SQL: UPDATE posts SET deleted_at = ? WHERE id = ?"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 6, "PostService", "PostController", "処理完了", "-"],
        ["SHOULD", "S2", "投稿論理削除 (POST /posts/{id}/delete)", 7, "PostController", "ユーザー (ブラウザ)", "HTTP 302 Redirect (/posts)", "一覧画面へリダイレクト"],

        # SHOULD: S3: タグ機能 (タグ別一覧 GET /tags/{name}/posts)
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 1, "ユーザー (ブラウザ)", "TagController", "GET /tags/{name}/posts リクエスト送信", "-"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 2, "TagController", "PostService", "getPostsByTagName(name) 呼び出し", "-"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 3, "PostService", "PostRepository", "findByTagNameOrderByCreatedAtDesc()", "SQL: SELECT p.* FROM posts p JOIN post_tags pt ON p.id = pt.post_id JOIN tags t ON t.id = pt.tag_id WHERE t.name = ? AND p.deleted_at IS NULL ORDER BY p.created_at DESC"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 4, "PostRepository", "PostService", "投稿リスト返却 (List<Post>)", "-"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 5, "PostService", "TagController", "投稿リスト返却", "-"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 6, "TagController", "Model", "posts, tagName 属性を設定", "-"],
        ["SHOULD", "S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 7, "TagController", "Thymeleaf (tags/list.html)", "テンプレート描画・HTML返却", "指定したタグに紐づく有効な投稿一覧を表示"],

        # COULD: C1: アバターカラー選択
        ["COULD", "C1", "アバターカラー選択 (GET/POST)", 1, "ユーザー (ブラウザ)", "PostController", "新規作成画面(GET /posts/new)を表示", "10色のカラーパレットチップから選択可能"],
        ["COULD", "C1", "アバターカラー選択 (GET/POST)", 2, "ユーザー (ブラウザ)", "PostController", "POST /posts でカラーコード値を送信", "例: '#3B82F6'"],
        ["COULD", "C1", "アバターカラー選択 (GET/POST)", 3, "PostController", "PostService", "Postエンティティに color を格納して保存", "SQL: INSERT / UPDATE 時の color カラム値設定"],

        # COULD: C2: ごみ箱機能（一覧・復元・空にする）
        ["COULD", "C2", "ゴミ箱一覧 (GET /posts/trash)", 1, "ユーザー (ブラウザ)", "PostController", "GET /posts/trash リクエスト送信", "-"],
        ["COULD", "C2", "ゴミ箱一覧 (GET /posts/trash)", 2, "PostController", "PostService", "getTrashPosts() 呼び出し", "-"],
        ["COULD", "C2", "ゴミ箱一覧 (GET /posts/trash)", 3, "PostService", "PostRepository", "findByDeletedAtIsNotNullAndPurgedAtIsNull()", "SQL: SELECT * FROM posts WHERE deleted_at IS NOT NULL AND purged_at IS NULL ORDER BY deleted_at DESC"],
        ["COULD", "C2", "ゴミ箱一覧 (GET /posts/trash)", 4, "PostController", "Thymeleaf (posts/trash.html)", "テンプレート描画・HTML返却", "ゴミ箱に入っている投稿のみ表示。復元・完全削除ボタンを配置"],
        ["COULD", "C2", "ゴミ箱復元 (POST /posts/{id}/restore)", 5, "ユーザー (ブラウザ)", "PostController", "POST /posts/{id}/restore 送信", "CSRFトークン検証含む"],
        ["COULD", "C2", "ゴミ箱復元 (POST /posts/{id}/restore)", 6, "PostController", "PostService", "restorePost(id) 呼び出し", "-"],
        ["COULD", "C2", "ゴミ箱復元 (POST /posts/{id}/restore)", 7, "PostService", "Post", "restore() 呼び出し", "エンティティの deletedAt を NULL に設定"],
        ["COULD", "C2", "ゴミ箱復元 (POST /posts/{id}/restore)", 8, "PostService", "PostRepository", "save(post) 呼び出し", "SQL: UPDATE posts SET deleted_at = NULL WHERE id = ?"],
        ["COULD", "C2", "ゴミ箱を空にする (POST /posts/trash/empty)", 9, "ユーザー (ブラウザ)", "PostController", "POST /posts/trash/empty 送信", "-"],
        ["COULD", "C2", "ゴミ箱を空にする (POST /posts/trash/empty)", 10, "PostController", "PostService", "emptyTrash() 呼び出し", "-"],
        ["COULD", "C2", "ゴミ箱を空にする (POST /posts/trash/empty)", 11, "PostService", "PostRepository", "findByDeletedAtIsNotNullAndPurgedAtIsNull()", "対象全件の取得"],
        ["COULD", "C2", "ゴミ箱を空にする (POST /posts/trash/empty)", 12, "PostService", "Post", "purge() 呼び出し", "対象全件に対して purgedAt に現在日時をセット"],
        ["COULD", "C2", "ゴミ箱を空にする (POST /posts/trash/empty)", 13, "PostService", "PostRepository", "save(post) 呼び出し", "SQL: UPDATE posts SET purged_at = ? WHERE id = ?"],

        # COULD: C3: REST API & API仕様書
        ["COULD", "C3", "REST API (GET /api/posts)", 1, "外部クライアント", "PostApiController", "GET /api/posts?all=true 送信", "セキュリティ設定でCSRF除外済み"],
        ["COULD", "C3", "REST API (GET /api/posts)", 2, "PostApiController", "PostRepository", "findTop50ByOrderByCreatedAtDesc() (all=true)", "SQL: SELECT * FROM posts ORDER BY created_at DESC (all=false時は有効な投稿のみ取得)"],
        ["COULD", "C3", "REST API (GET /api/posts)", 3, "PostApiController", "PostResponse (DTO)", "PostResponse(post) インスタンス生成", "deleteStatus (0:有効, 1:ゴミ箱, 2:完全削除) および 各種日付(ISO-8601)・editedAt のマッピング"],
        ["COULD", "C3", "REST API (GET /api/posts)", 4, "PostApiController", "外部クライアント", "JSON レスポンス返却 (List<PostResponse>)", "-"],
        ["COULD", "C3", "API仕様書配信 (GET /api-docs)", 5, "外部クライアント / ブラウザ", "OpenApiController", "GET /api-docs リクエスト送信", "-"],
        ["COULD", "C3", "API仕様書配信 (GET /api-docs)", 6, "OpenApiController", "Resource", "openapi.json 読み込み", "-"],
        ["COULD", "C3", "API仕様書配信 (GET /api-docs)", 7, "OpenApiController", "外部クライアント", "JSON 形式で仕様書テキストをそのまま配信", "-"],

        # 追加: 投稿編集機能
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 1, "ユーザー (ブラウザ)", "PostController", "GET /posts/{id}/edit リクエスト送信", "-"],
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 2, "PostController", "PostService", "findById(id) 呼び出し", "投稿の取得。論理削除済みは404"],
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 3, "PostController", "ClientHashGenerator", "generate(ip, userAgent) 呼び出し", "閲覧者のクライアントハッシュを生成"],
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 4, "PostController", "PostController", "本人検証", "post.clientHash != currentHash の場合は AccessDeniedException (403)"],
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 5, "PostController", "Model", "既存データ(本文, カラー, タグカンマ区切り文字列)をPostFormに設定して追加", "-"],
        ["追加", "Edit", "投稿編集画面表示 (GET /posts/{id}/edit)", 6, "PostController", "Thymeleaf (posts/edit.html)", "テンプレート描画・HTML返却", "投稿者名は readonly 表示、本文・カラーチップ・タグ入力欄を配置"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 7, "ユーザー (ブラウザ)", "PostController", "POST /posts/{id}/edit 送信", "CSRFトークン検証、入力値バリデーション"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 8, "PostController", "ClientHashGenerator", "generate(ip, userAgent)", "送信者のクライアントハッシュを生成"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 9, "PostController", "PostService", "updatePost(id, form, hash) 呼び出し", "-"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 10, "PostService", "PostRepository", "findById(id)", "存在確認"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 11, "PostService", "PostService", "本人検証", "post.clientHash != hash の場合は AccessDeniedException (403)"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 12, "PostService", "Post", "update(body, color) 呼び出し", "エンティティの body, color を更新し、editedAt に現在日時をセット"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 13, "PostService", "TagRepository & Post", "タグ of 差分紐付け更新・クリーンアップ", "解除されたタグが他で使われていなければ tagRepository.delete で物理削除。新規タグを追加"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 14, "PostService", "PostRepository", "save(post)", "SQL: UPDATE posts SET body = ?, color = ?, edited_at = ? WHERE id = ?"],
        ["追加", "Edit", "投稿編集保存 (POST /posts/{id}/edit)", 15, "PostController", "ユーザー (ブラウザ)", "HTTP 302 Redirect (/posts/{id})", "更新成功後、詳細画面へリダイレクト"]
    ]
    
    # 書き込み
    ws1.append(headers)
    for row in data:
        ws1.append(row)
        
    font_family = "Segoe UI"
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    
    fill_must = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_should = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_could = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_extra = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    
    for row_idx in range(2, len(data) + 2):
        category = ws1.cell(row=row_idx, column=1).value
        
        if category == "MUST":
            cat_fill = fill_must
        elif category == "SHOULD":
            cat_fill = fill_should
        elif category == "COULD":
            cat_fill = fill_could
        else:
            cat_fill = fill_extra
            
        for col_idx in range(1, len(headers) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.font = Font(name=font_family, size=10)
            cell.border = border_thin
            
            if col_idx in [1, 2]:
                cell.fill = cat_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
                
            if col_idx in [7, 8]:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                
    ws1.row_dimensions[1].height = 28
    for row_idx in range(2, len(data) + 2):
        ws1.row_dimensions[row_idx].height = 24
        
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            val_len = sum(2 if ord(char) > 255 else 1 for char in val)
            if val_len > max_len:
                max_len = val_len
        ws1.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 60)

    # -------------------------------------------------------------
    # シート2: Mermaid Diagrams + D列に画像貼り付け
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Mermaid Diagrams")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.append(["機能ID", "機能名", "Mermaid シーケンスコード (コピペ用)", "レンダリングイメージ図"])
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=1).font = header_font
    ws2.cell(row=1, column=2).fill = header_fill
    ws2.cell(row=1, column=2).font = header_font
    ws2.cell(row=1, column=3).fill = header_fill
    ws2.cell(row=1, column=3).font = header_font
    ws2.cell(row=1, column=4).fill = header_fill
    ws2.cell(row=1, column=4).font = header_font
    
    mermaid_data = [
        ("M1", "投稿一覧表示 (GET /posts)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n"
         "    participant DB as データベース\n\n"
         "    Browser->>Ctrl: GET /posts\n"
         "    Ctrl->>Service: getPosts()\n"
         "    Service->>Repo: findTop50ByDeletedAtIsNullOrderByCreatedAtDesc()\n"
         "    Repo->>DB: SELECT * FROM posts WHERE deleted_at IS NULL ORDER BY created_at DESC (最大50件)\n"
         "    DB-->>Repo: 投稿データ\n"
         "    Repo-->>Service: List<Post>\n"
         "    Service-->>Ctrl: List<Post>\n"
         "    Ctrl->>Ctrl: Modelに 'posts' を設定\n"
         "    Ctrl-->>Browser: Thymeleaf テンプレート (posts/list) 描画・HTML返却"),
         
        ("M2", "新規投稿フォーム表示 (GET /posts/new)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n\n"
         "    Browser->>Ctrl: GET /posts/new\n"
         "    Ctrl->>Ctrl: Modelに新規 PostForm を設定\n"
         "    Ctrl-->>Browser: Thymeleaf テンプレート (posts/form) 描画・HTML返却"),
         
        ("M3", "投稿登録 (POST /posts)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant HashGen as ClientHashGenerator\n"
         "    participant Service as PostService\n"
         "    participant TagRepo as TagRepository\n"
         "    participant Repo as PostRepository\n\n"
         "    Browser->>Ctrl: POST /posts (Formデータ + CSRF)\n"
         "    Ctrl->>Ctrl: バリデーション\n"
         "    alt エラーあり\n"
         "        Ctrl-->>Browser: posts/form 再表示 (200)\n"
         "    else エラーなし\n"
         "        Ctrl->>HashGen: generate(IP, UA)\n"
         "        HashGen-->>Ctrl: clientHash (8桁)\n"
         "        Ctrl->>Service: create(form, clientHash)\n"
         "        Service->>TagRepo: findByName(tagName) / save(tag)\n"
         "        Service->>Repo: save(post) (INSERT)\n"
         "        Repo-->>Service: 保存済み Post\n"
         "        Service-->>Ctrl: 処理完了\n"
         "        Ctrl-->>Browser: HTTP 302 Redirect (/posts)\n"
         "    end"),
         
        ("M4", "投稿詳細表示 (GET /posts/{id})", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n"
         "    participant LikeService as PostLikeService\n"
         "    participant HashGen as ClientHashGenerator\n\n"
         "    Browser->>Ctrl: GET /posts/{id}\n"
         "    Ctrl->>Service: findById(id)\n"
         "    Service->>Repo: findById(id)\n"
         "    Repo-->>Service: Postエンティティ\n"
         "    alt 削除済み\n"
         "        Service-->>Ctrl: PostNotFoundException (404)\n"
         "    else 有効な投稿\n"
         "        Service-->>Ctrl: Post\n"
         "    end\n"
         "    Ctrl->>LikeService: countByPostId(id)\n"
         "    LikeService-->>Ctrl: いいね件数\n"
         "    Ctrl->>HashGen: generate(IP, UA)\n"
         "    HashGen-->>Ctrl: currentHash\n"
         "    Ctrl->>Ctrl: Model設定\n"
         "    Ctrl-->>Browser: Thymeleaf テンプレート (posts/detail) 描画・HTML返却"),
         
        ("S1", "いいねトグル (POST /posts/{id}/likes)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant HashGen as ClientHashGenerator\n"
         "    participant LikeService as PostLikeService\n"
         "    participant Repo as PostRepository\n"
         "    participant LikeRepo as PostLikeRepository\n\n"
         "    Browser->>Ctrl: POST /posts/{id}/likes (CSRF付き)\n"
         "    Ctrl->>HashGen: generate(IP, UA)\n"
         "    HashGen-->>Ctrl: clientHash\n"
         "    Ctrl->>LikeService: toggle(id, clientHash)\n"
         "    LikeService->>Repo: findById(id)\n"
         "    LikeService->>LikeRepo: findByPostIdAndClientHash()\n"
         "    alt 登録済み\n"
         "        LikeService->>LikeRepo: delete(like) (いいね解除)\n"
         "    else 未登録\n"
         "        LikeService->>LikeRepo: save(new PostLike) (いいね登録)\n"
         "    end\n"
         "    LikeService-->>Ctrl: 処理完了\n"
         "    Ctrl-->>Browser: HTTP 302 Redirect (/posts/{id})"),
         
        ("S2", "投稿論理削除 (POST /posts/{id}/delete)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n\n"
         "    Browser->>Ctrl: POST /posts/{id}/delete (CSRF付き)\n"
         "    Ctrl->>Service: deletePost(id)\n"
         "    Service->>Repo: findById(id)\n"
         "    Service->>Repo: save(post) (deleted_at設定)\n"
         "    Service-->>Ctrl: 処理完了\n"
         "    Ctrl-->>Browser: HTTP 302 Redirect (/posts)"),
         
        ("S3", "タグ別投稿一覧表示 (GET /tags/{name}/posts)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as TagController\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n\n"
         "    Browser->>Ctrl: GET /tags/{name}/posts\n"
         "    Ctrl->>Service: getPostsByTagName(name)\n"
         "    Service->>Repo: findByTagNameOrderByCreatedAtDesc()\n"
         "    Repo-->>Service: List<Post> (有効な投稿一覧)\n"
         "    Service-->>Ctrl: List<Post>\n"
         "    Ctrl->>Ctrl: Model設定\n"
         "    Ctrl-->>Browser: Thymeleaf テンプレート (tags/list) 描画・HTML返却"),
         
        ("C2", "投稿ごみ箱機能（一覧・復元・空にする）", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n\n"
         "    Note over Browser, Repo: 【ゴミ箱一覧表示】\n"
         "    Browser->>Ctrl: GET /posts/trash\n"
         "    Ctrl->>Service: getTrashPosts()\n"
         "    Service->>Repo: findByDeletedAtIsNotNullAndPurgedAtIsNull()\n"
         "    Repo-->>Service: List<Post>\n"
         "    Service-->>Ctrl: List<Post>\n"
         "    Ctrl-->>Browser: posts/trash.html 描画\n\n"
         "    Note over Browser, Repo: 【投稿の復元】\n"
         "    Browser->>Ctrl: POST /posts/{id}/restore\n"
         "    Ctrl->>Service: restorePost(id)\n"
         "    Service->>Repo: findById(id)\n"
         "    Service->>Repo: save(post) (deleted_at を NULL)\n"
         "    Service-->>Ctrl: 処理完了\n"
         "    Ctrl-->>Browser: HTTP 302 Redirect (/posts/trash)\n\n"
         "    Note over Browser, Repo: 【ゴミ箱を空にする】\n"
         "    Browser->>Ctrl: POST /posts/trash/empty\n"
         "    Ctrl->>Service: emptyTrash()\n"
         "    Service->>Repo: findByDeletedAtIsNotNullAndPurgedAtIsNull()\n"
         "    Repo-->>Service: 削除済み投稿リスト\n"
         "    loop 各投稿に対して\n"
         "        Service->>Repo: save(post) (purged_at 設定)\n"
         "    end\n"
         "    Service-->>Ctrl: 処理完了\n"
         "    Ctrl-->>Browser: HTTP 302 Redirect (/posts/trash)"),
         
        ("C3", "REST API & API仕様書 (GET /api/posts, GET /api-docs)", 
         "sequenceDiagram\n"
         "    actor Client as 外部クライアント\n"
         "    participant Ctrl as PostApiController\n"
         "    participant OpenCtrl as OpenApiController\n"
         "    participant Repo as PostRepository\n\n"
         "    Note over Client, Repo: 【REST API 投稿取得】\n"
         "    Client->>Ctrl: GET /api/posts?all=true\n"
         "    Ctrl->>Repo: findTop50ByOrderByCreatedAtDesc() (all=true)\n"
         "    Repo-->>Ctrl: List<Post>\n"
         "    Ctrl->>Ctrl: PostResponseへマッピング (editedAt等含む)\n"
         "    Ctrl-->>Client: JSONレスポンス返却 (HTTP 200)\n\n"
         "    Note over Client, Repo: 【API仕様書取得】\n"
         "    Client->>OpenCtrl: GET /api-docs\n"
         "    OpenCtrl->>OpenCtrl: openapi.json 読み込み\n"
         "    OpenCtrl-->>Client: JSON形式の仕様書返却 (HTTP 200)"),
         
        ("Edit", "投稿編集機能 (GET/POST /posts/{id}/edit)", 
         "sequenceDiagram\n"
         "    actor Browser as ユーザー (ブラウザ)\n"
         "    participant Ctrl as PostController\n"
         "    participant HashGen as ClientHashGenerator\n"
         "    participant Service as PostService\n"
         "    participant Repo as PostRepository\n"
         "    participant TagRepo as TagRepository\n\n"
         "    Note over Browser, TagRepo: 【編集画面表示】\n"
         "    Browser->>Ctrl: GET /posts/{id}/edit\n"
         "    Ctrl->>Service: findById(id)\n"
         "    Service-->>Ctrl: Post\n"
         "    Ctrl->>HashGen: generate(IP, UA)\n"
         "    HashGen-->>Ctrl: currentHash (閲覧者ハッシュ)\n"
         "    alt 本人ハッシュ不一致\n"
         "        Ctrl-->>Browser: AccessDeniedException (403)\n"
         "    else 一致\n"
         "        Ctrl->>Ctrl: Modelに既存データをバインド\n"
         "        Ctrl-->>Browser: posts/edit.html 返却 (200)\n"
         "    end\n\n"
         "    Note over Browser, TagRepo: 【編集内容の保存】\n"
         "    Browser->>Ctrl: POST /posts/{id}/edit (Formデータ + CSRF)\n"
         "    Ctrl->>Ctrl: バリデーション実行 (@Valid)\n"
         "    Ctrl->>HashGen: generate(IP, UA)\n"
         "    HashGen-->>Ctrl: clientHash\n"
         "    Ctrl->>Service: updatePost(id, form, clientHash)\n"
         "    Service->>Repo: findById(id)\n"
         "    Service->>Service: 本人ハッシュ検証 (不一致なら 403スロー)\n"
         "    Service->>Service: post.update(body, color) (edited_at設定)\n"
         "    Service->>TagRepo: タグの差分紐付け更新・浮いたタグ削除\n"
         "    Service->>Repo: save(post) (UPDATE)\n"
         "    Service-->>Ctrl: 処理完了\n"
         "    Ctrl-->>Browser: HTTP 302 Redirect (/posts/{id})")
    ]
    
    for idx, (fid, fname, code) in enumerate(mermaid_data):
        row_idx = idx + 2
        ws2.cell(row=row_idx, column=1, value=fid)
        ws2.cell(row=row_idx, column=2, value=fname)
        ws2.cell(row=row_idx, column=3, value=code)
        
        # 画像ダウンロード＆Excelに埋め込み
        img_path = generate_mermaid_image(code, fid)
        if img_path and os.path.exists(img_path):
            try:
                img = OpenpyxlImage(img_path)
                original_width = img.width
                original_height = img.height
                if original_height > 0:
                    ratio = 320 / original_height
                    img.width = int(original_width * ratio)
                    img.height = 320
                
                cell_address = f"D{row_idx}"
                ws2.add_image(img, cell_address)
            except Exception as ex:
                print(f"Error embedding image for {fid}: {ex}")
        else:
            ws2.cell(row=row_idx, column=4, value="画像取得失敗")
            
    for col_idx in [1, 2, 3, 4]:
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin
        
    for row_idx in range(2, len(mermaid_data) + 2):
        ws2.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=row_idx, column=2).alignment = Alignment(vertical="center", wrap_text=True)
        ws2.cell(row=row_idx, column=3).alignment = Alignment(vertical="top", wrap_text=True)
        ws2.cell(row=row_idx, column=4).alignment = Alignment(horizontal="center", vertical="center")
        
        ws2.row_dimensions[row_idx].height = 260
        
        for col_idx in [1, 2, 3, 4]:
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.font = Font(name=font_family, size=10)
            cell.border = border_thin
            
    ws2.row_dimensions[1].height = 28
    ws2.column_dimensions['A'].width = 12
    ws2.column_dimensions['B'].width = 30
    ws2.column_dimensions['C'].width = 50
    ws2.column_dimensions['D'].width = 65

    file_path = "g:\\ya-work\\workspace\\tsubuyaki-board\\scratch\\tsubuyaki_sequence_diagrams.xlsx"
    wb.save(file_path)
    print(f"Excel file generated successfully: {file_path}")

if __name__ == "__main__":
    create_sequence_excel()
